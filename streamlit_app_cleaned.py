
import streamlit as st
import streamlit.components.v1 as components
from docx import Document
import re
import uuid
import os
import time
import html
import csv
from io import BytesIO

# --------------------------- إعدادات الصفحة ----------------------------
st.set_page_config(
    page_title="القوانين اليمنية بآخر تعديلاتها حتى عام 2025م",
    layout="wide",
    initial_sidebar_state="expanded"
)

# --------------------------- CSS لتنسيق RTL ----------------------------
components.html("""
<style>
textarea, input[type="text"], .stTextArea textarea, .stTextInput input {
    direction: rtl !important;
    text-align: right !important;
    font-family: "Tahoma", "Arial", sans-serif !important;
    font-size: 18px !important;
}
.scroll-btn {
    position: fixed;
    left: 10px;
    padding: 12px;
    font-size: 24px;
    border-radius: 50%;
    background-color: #c5e1a5;
    color: black;
    cursor: pointer;
    z-index: 9999;
    border: none;
    box-shadow: 1px 1px 5px #888;
}
#scroll-top-btn { bottom: 80px; }
#scroll-bottom-btn { bottom: 20px; }
.rtl-metric {
    direction: rtl;
    text-align: right !important;
    margin-right: 0 !important;
}
.rtl-metric .stMetricDelta, .rtl-download-btn, .stButton, .stDownloadButton, .stMetric {
    direction: rtl !important;
    text-align: right !important;
}
</style>
<button class='scroll-btn' id='scroll-top-btn' onclick='window.scrollTo({top: 0, behavior: "smooth"});'>⬆️</button>
<button class='scroll-btn' id='scroll-bottom-btn' onclick='window.scrollTo({top: document.body.scrollHeight, behavior: "smooth"});'>⬇️</button>
""", height=1)

# --------------------------- الثوابت ----------------------------
TRIAL_DURATION = 3 * 24 * 60 * 60
TRIAL_USERS_FILE = "trial_users.txt"
DEVICE_ID_FILE = "device_id.txt"
ACTIVATED_FILE = "activated.txt"
ACTIVATION_CODES_FILE = "activation_codes.txt"
LAWS_DIR = "laws"

# --------------------------- دوال مساعدة ----------------------------
def rtl_label(text):
    st.markdown(f'<div style="direction: rtl; text-align: right;">{text}</div>', unsafe_allow_html=True)

def get_device_id():
    if os.path.exists(DEVICE_ID_FILE):
        with open(DEVICE_ID_FILE, "r") as f:
            return f.read().strip()
    new_id = str(uuid.uuid4())
    with open(DEVICE_ID_FILE, "w") as f:
        f.write(new_id)
    return new_id

def get_trial_start(device_id):
    if not os.path.exists(TRIAL_USERS_FILE):
        return None
    with open(TRIAL_USERS_FILE, "r") as f:
        reader = csv.reader(f)
        for row in reader:
            if row and row[0] == device_id:
                return float(row[1])
    return None

def register_trial(device_id):
    if not os.path.exists(TRIAL_USERS_FILE):
        with open(TRIAL_USERS_FILE, "w", newline='') as f:
            pass
    with open(TRIAL_USERS_FILE, "a", newline='') as f:
        writer = csv.writer(f)
        writer.writerow([device_id, time.time()])

def is_activated():
    return os.path.exists(ACTIVATED_FILE)

def activate_app(code):
    if not os.path.exists(ACTIVATION_CODES_FILE):
        return False
    with open(ACTIVATION_CODES_FILE, "r") as f:
        codes = [line.strip() for line in f.readlines()]
    if code in codes:
        codes.remove(code)
        with open(ACTIVATION_CODES_FILE, "w") as f:
            for c in codes:
                f.write(c + "\n")
        with open(ACTIVATED_FILE, "w") as f:
            f.write("activated")
        return True
    return False

def normalize_arabic_numbers(text):
    return text.translate(str.maketrans('٠١٢٣٤٥٦٧٨٩', '0123456789'))

def highlight_keywords(text, keywords):
    for kw in keywords:
        text = re.sub(f"({re.escape(kw)})", r"<mark>\1</mark>", text, flags=re.IGNORECASE)
    return text

def export_results_to_word(results):
    document = Document()
    document.add_heading('نتائج البحث في القوانين اليمنية', level=1)
    if not results:
        document.add_paragraph("لم يتم العثور على نتائج.")
    else:
        for i, r in enumerate(results):
            document.add_heading(f"القانون: {r['law']} - المادة: {r['num']}", level=2)
            document.add_paragraph(r['plain'])
            if i < len(results) - 1:
                document.add_page_break()
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

# --------------------------- الدالة الرئيسية (placeholder) ----------------------------
if __name__ == "__main__":
    st.title("نظام بحث القوانين اليمنية")
    rtl_label("يرجى تفعيل التطبيق أو بدء النسخة التجريبية")
